"""Regression tests for invoice uploader/date filters and XLSX export."""
import io
import zipfile
from datetime import datetime, timezone as dt_timezone
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from django.conf import settings
from rest_framework.test import APIClient

from apps.invoices.models import Invoice, InvoiceStatus
from apps.users.models import User, UserRole


@pytest.fixture
def admin(db):
    return User.objects.create_user(
        username='invoice_filter_admin',
        password='testpass123',
        role=UserRole.SUPER_ADMIN,
        is_active=True,
    )


@pytest.fixture
def uploader(db):
    return User.objects.create_user(
        username='invoice_uploader',
        password='testpass123',
        role=UserRole.ITEM_MANAGER,
        is_active=True,
    )


def create_invoice(*, creator, name, created_at, amount='10.00'):
    invoice = Invoice.objects.create(
        name=name,
        amount=Decimal(amount),
        category='采购',
        custom_type='测试',
        project='项目A',
        file=f'invoices/2099/01/{name}.pdf',
        original_filename=f'{name}.pdf',
        status=InvoiceStatus.PENDING,
        item=None,
        registered_month='2099-01',
        created_by=creator,
    )
    Invoice.objects.filter(pk=invoice.pk).update(created_at=created_at)
    invoice.refresh_from_db()
    return invoice


def materialize(invoice):
    path = Path(settings.MEDIA_ROOT) / invoice.file
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b'%PDF-1.4 filtered export')


@pytest.mark.django_db
class TestInvoiceFilteredExports:
    def test_list_filters_by_uploader_and_custom_date_range(self, admin, uploader):
        in_range = create_invoice(
            creator=uploader,
            name='uploader-in-range',
            created_at=datetime(2026, 8, 10, 3, 0, tzinfo=dt_timezone.utc),
        )
        create_invoice(
            creator=uploader,
            name='uploader-out-of-range',
            created_at=datetime(2026, 7, 1, 3, 0, tzinfo=dt_timezone.utc),
        )
        create_invoice(
            creator=admin,
            name='other-uploader',
            created_at=datetime(2026, 8, 10, 3, 0, tzinfo=dt_timezone.utc),
        )
        client = APIClient()
        client.force_authenticate(user=admin)

        response = client.get(
            f'/api/v1/invoices/?uploader={uploader.id}&date_from=2026-08-01&date_to=2026-08-31&page_size=100'
        )

        assert response.status_code == 200
        assert [row['id'] for row in response.data['results']] == [in_range.id]

    def test_uploaders_endpoint_returns_users_with_invoices(self, admin, uploader):
        create_invoice(
            creator=uploader,
            name='has-uploader',
            created_at=datetime(2026, 8, 10, 3, 0, tzinfo=dt_timezone.utc),
        )
        client = APIClient()
        client.force_authenticate(user=admin)

        response = client.get('/api/v1/invoices/uploaders/')

        assert response.status_code == 200
        assert {'id': uploader.id, 'username': uploader.username} in response.data

    def test_reimbursement_excel_respects_uploader_and_date_range(self, admin, uploader):
        in_range = create_invoice(
            creator=uploader,
            name='excel-in-range',
            created_at=datetime(2026, 8, 10, 3, 0, tzinfo=dt_timezone.utc),
            amount='88.50',
        )
        create_invoice(
            creator=admin,
            name='excel-other-uploader',
            created_at=datetime(2026, 8, 10, 3, 0, tzinfo=dt_timezone.utc),
        )
        client = APIClient()
        client.force_authenticate(user=admin)

        response = client.get(
            f'/api/v1/invoices/export/reimbursement/?uploader={uploader.id}'
            '&date_from=2026-08-01&date_to=2026-08-31'
        )

        assert response.status_code == 200
        assert response['Content-Type'].startswith(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = workbook['报销清单']
        assert [cell.value for cell in sheet[1]][:11] == [
            '序号', '发票名称', '金额（元）', '类别', '类型', '项目',
            '报销状态', '关联物品', '登记月份', '上传者', '上传时间',
        ]
        assert sheet.cell(2, 1).value == 1
        assert sheet.cell(2, 2).value == in_range.name
        assert sheet.cell(2, 3).value == 88.5
        assert sheet.cell(2, 10).value == uploader.username
        assert sheet.cell(3, 2).value == '合计'
        assert sheet.cell(3, 3).value == 88.5
        assert sheet.max_row == 3

    def test_zip_export_respects_uploader_filter(self, admin, uploader):
        wanted = create_invoice(
            creator=uploader,
            name='zip-wanted',
            created_at=datetime(2026, 8, 10, 3, 0, tzinfo=dt_timezone.utc),
        )
        unwanted = create_invoice(
            creator=admin,
            name='zip-unwanted',
            created_at=datetime(2026, 8, 10, 3, 0, tzinfo=dt_timezone.utc),
        )
        materialize(wanted)
        materialize(unwanted)
        client = APIClient()
        client.force_authenticate(user=admin)

        response = client.get(f'/api/v1/invoices/export/?uploader={uploader.id}')

        assert response.status_code == 200
        assert response['Content-Type'] == 'application/zip'
        with zipfile.ZipFile(io.BytesIO(response.content), 'r') as archive:
            names = archive.namelist()
        assert len(names) == 1
        assert 'zip-wanted' in names[0]
        assert 'zip-unwanted' not in names[0]

    def test_invalid_date_range_returns_400(self, admin):
        client = APIClient()
        client.force_authenticate(user=admin)

        response = client.get(
            '/api/v1/invoices/export/reimbursement/?date_from=2026-09-01&date_to=2026-08-01'
        )

        assert response.status_code == 400
        assert response.data['code'] == 'INVALID_FILTER'
