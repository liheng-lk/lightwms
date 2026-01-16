"""
Import service for LightWMS.

Handles CSV/XLSX file parsing, validation, and item import with SKU-based upsert.
"""
import csv
import io
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

from django.db import transaction

from apps.inventory.models import (
    Item, ItemType, ItemStatus, Location, Category, Warehouse
)
from apps.users.models import User
from apps.audit.models import AuditLog, AuditAction, TargetType
from core.exceptions import BusinessConflictError


@dataclass
class ImportError:
    """Represents a single import error."""
    row: int
    field: str
    message: str


@dataclass
class ImportResult:
    """Result of an import operation."""
    success_count: int = 0
    update_count: int = 0
    create_count: int = 0
    error_count: int = 0
    errors: List[ImportError] = field(default_factory=list)
    
    @property
    def total_processed(self) -> int:
        return self.success_count + self.error_count


class ImportService:
    """
    Service for importing items from CSV/XLSX files.
    
    Features:
    - SKU-based upsert (create if not exists, update if exists)
    - Per-row validation with detailed error reporting
    - File size limit (10MB) and row limit (10000)
    - Support for CSV and XLSX formats
    """
    
    # Limits
    MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
    MAX_ROWS = 10000
    
    # Required and optional import fields
    REQUIRED_FIELDS = ['name', 'sku', 'type', 'warehouse_code', 'location_code']
    OPTIONAL_FIELDS = [
        'barcode', 'quantity', 'unit', 'description',
        'category_path', 'owner_username_or_phone'
    ]
    ALL_FIELDS = REQUIRED_FIELDS + OPTIONAL_FIELDS
    
    # Valid item types
    VALID_TYPES = [ItemType.REUSABLE, ItemType.CONSUMABLE]
    
    @classmethod
    def parse_csv(cls, file_content: bytes) -> Tuple[List[str], List[Dict[str, str]]]:
        """
        Parse CSV file content.
        
        Args:
            file_content: Raw bytes of CSV file
            
        Returns:
            Tuple of (headers, rows) where rows is list of dicts
            
        Raises:
            BusinessConflictError: If file is invalid
        """
        try:
            # Try UTF-8 first, then fall back to latin-1
            try:
                text = file_content.decode('utf-8-sig')  # Handle BOM
            except UnicodeDecodeError:
                text = file_content.decode('latin-1')
            
            reader = csv.DictReader(io.StringIO(text))
            headers = reader.fieldnames or []
            rows = list(reader)
            
            return headers, rows
        except Exception as e:
            raise BusinessConflictError(
                code='INVALID_FILE',
                message=f'Failed to parse CSV file: {str(e)}',
                details={'error': str(e)}
            )
    
    @classmethod
    def parse_xlsx(cls, file_content: bytes) -> Tuple[List[str], List[Dict[str, str]]]:
        """
        Parse XLSX file content.
        
        Args:
            file_content: Raw bytes of XLSX file
            
        Returns:
            Tuple of (headers, rows) where rows is list of dicts
            
        Raises:
            BusinessConflictError: If file is invalid
        """
        try:
            from openpyxl import load_workbook
            
            workbook = load_workbook(filename=io.BytesIO(file_content), read_only=True)
            sheet = workbook.active
            
            rows_iter = sheet.iter_rows(values_only=True)
            
            # First row is headers
            header_row = next(rows_iter, None)
            if not header_row:
                return [], []
            
            headers = [str(h).strip() if h else '' for h in header_row]
            
            # Parse data rows
            rows = []
            for row in rows_iter:
                row_dict = {}
                for i, value in enumerate(row):
                    if i < len(headers) and headers[i]:
                        # Convert to string, handle None
                        row_dict[headers[i]] = str(value).strip() if value is not None else ''
                rows.append(row_dict)
            
            workbook.close()
            return headers, rows
        except Exception as e:
            raise BusinessConflictError(
                code='INVALID_FILE',
                message=f'Failed to parse XLSX file: {str(e)}',
                details={'error': str(e)}
            )

    @classmethod
    def validate_file_size(cls, file_size: int) -> None:
        """
        Validate file size is within limits.
        
        Args:
            file_size: Size of file in bytes
            
        Raises:
            BusinessConflictError: If file is too large
        """
        if file_size > cls.MAX_FILE_SIZE:
            raise BusinessConflictError(
                code='FILE_TOO_LARGE',
                message=f'File size exceeds maximum of {cls.MAX_FILE_SIZE // (1024*1024)}MB.',
                details={'max_size': cls.MAX_FILE_SIZE, 'actual_size': file_size}
            )
    
    @classmethod
    def validate_row_count(cls, row_count: int) -> None:
        """
        Validate row count is within limits.
        
        Args:
            row_count: Number of data rows
            
        Raises:
            BusinessConflictError: If too many rows
        """
        if row_count > cls.MAX_ROWS:
            raise BusinessConflictError(
                code='ROW_LIMIT_EXCEEDED',
                message=f'Row count exceeds maximum of {cls.MAX_ROWS}.',
                details={'max_rows': cls.MAX_ROWS, 'actual_rows': row_count}
            )
    
    @classmethod
    def validate_headers(cls, headers: List[str]) -> List[ImportError]:
        """
        Validate that required headers are present.
        
        Args:
            headers: List of header names from file
            
        Returns:
            List of errors for missing headers
        """
        errors = []
        headers_lower = [h.lower() for h in headers]
        
        for required in cls.REQUIRED_FIELDS:
            if required.lower() not in headers_lower:
                errors.append(ImportError(
                    row=0,
                    field=required,
                    message=f'Required column "{required}" is missing.'
                ))
        
        return errors
    
    @classmethod
    def _normalize_headers(cls, headers: List[str]) -> Dict[str, str]:
        """
        Create mapping from lowercase header to original header.
        
        Args:
            headers: Original headers
            
        Returns:
            Dict mapping lowercase to original
        """
        return {h.lower(): h for h in headers}
    
    @classmethod
    def _get_field_value(
        cls,
        row: Dict[str, str],
        field: str,
        header_map: Dict[str, str]
    ) -> str:
        """
        Get field value from row, handling case-insensitive lookup.
        
        Args:
            row: Row data dict
            field: Field name to look up
            header_map: Mapping from lowercase to original header
            
        Returns:
            Field value or empty string
        """
        original_header = header_map.get(field.lower())
        if original_header and original_header in row:
            return row[original_header].strip()
        return ''
    
    @classmethod
    def validate_rows(
        cls,
        rows: List[Dict[str, str]],
        headers: List[str]
    ) -> Tuple[List[Dict[str, Any]], List[ImportError]]:
        """
        Validate all rows and collect errors.
        
        Args:
            rows: List of row dicts from parsed file
            headers: List of headers
            
        Returns:
            Tuple of (validated_rows, errors)
        """
        validated_rows = []
        errors = []
        header_map = cls._normalize_headers(headers)
        
        # Pre-fetch lookup data for validation
        warehouses = {w.code.upper(): w for w in Warehouse.objects.filter(is_active=True)}
        locations_by_warehouse = {}
        for loc in Location.objects.filter(is_active=True).select_related('warehouse'):
            wh_code = loc.warehouse.code.upper()
            if wh_code not in locations_by_warehouse:
                locations_by_warehouse[wh_code] = {}
            locations_by_warehouse[wh_code][loc.code.upper()] = loc
        
        categories = cls._build_category_lookup()
        users = cls._build_user_lookup()
        
        for row_num, row in enumerate(rows, start=2):  # Start at 2 (1 is header)
            row_errors = []
            validated_data = {}
            
            # Validate required fields
            name = cls._get_field_value(row, 'name', header_map)
            if not name:
                row_errors.append(ImportError(row=row_num, field='name', message='Name is required.'))
            else:
                validated_data['name'] = name
            
            sku = cls._get_field_value(row, 'sku', header_map)
            if not sku:
                row_errors.append(ImportError(row=row_num, field='sku', message='SKU is required.'))
            else:
                validated_data['sku'] = sku.upper()
            
            item_type = cls._get_field_value(row, 'type', header_map).lower()
            if not item_type:
                row_errors.append(ImportError(row=row_num, field='type', message='Type is required.'))
            elif item_type not in ['reusable', 'consumable']:
                row_errors.append(ImportError(
                    row=row_num, field='type',
                    message=f'Type must be "reusable" or "consumable", got "{item_type}".'
                ))
            else:
                validated_data['type'] = item_type
            
            # Validate warehouse and location
            warehouse_code = cls._get_field_value(row, 'warehouse_code', header_map).upper()
            location_code = cls._get_field_value(row, 'location_code', header_map).upper()
            
            if not warehouse_code:
                row_errors.append(ImportError(row=row_num, field='warehouse_code', message='Warehouse code is required.'))
            elif warehouse_code not in warehouses:
                row_errors.append(ImportError(
                    row=row_num, field='warehouse_code',
                    message=f'Warehouse "{warehouse_code}" not found or inactive.'
                ))
            else:
                validated_data['warehouse'] = warehouses[warehouse_code]
            
            if not location_code:
                row_errors.append(ImportError(row=row_num, field='location_code', message='Location code is required.'))
            elif warehouse_code in locations_by_warehouse:
                wh_locations = locations_by_warehouse[warehouse_code]
                if location_code not in wh_locations:
                    row_errors.append(ImportError(
                        row=row_num, field='location_code',
                        message=f'Location "{location_code}" not found in warehouse "{warehouse_code}".'
                    ))
                else:
                    validated_data['location'] = wh_locations[location_code]
            
            # Validate optional fields
            barcode = cls._get_field_value(row, 'barcode', header_map)
            if barcode:
                validated_data['barcode'] = barcode
            
            quantity_str = cls._get_field_value(row, 'quantity', header_map)
            if quantity_str:
                try:
                    quantity = int(quantity_str)
                    if quantity < 0:
                        row_errors.append(ImportError(
                            row=row_num, field='quantity',
                            message='Quantity must be non-negative.'
                        ))
                    else:
                        validated_data['quantity'] = quantity
                except ValueError:
                    row_errors.append(ImportError(
                        row=row_num, field='quantity',
                        message=f'Invalid quantity "{quantity_str}".'
                    ))
            
            unit = cls._get_field_value(row, 'unit', header_map)
            if unit:
                validated_data['unit'] = unit
            
            description = cls._get_field_value(row, 'description', header_map)
            if description:
                validated_data['description'] = description
            
            # Validate category path
            category_path = cls._get_field_value(row, 'category_path', header_map)
            if category_path:
                category = categories.get(category_path.lower())
                if category:
                    validated_data['category'] = category
                else:
                    row_errors.append(ImportError(
                        row=row_num, field='category_path',
                        message=f'Category "{category_path}" not found.'
                    ))
            
            # Validate owner
            owner_ref = cls._get_field_value(row, 'owner_username_or_phone', header_map)
            if owner_ref:
                owner = users.get(owner_ref.lower())
                if owner:
                    validated_data['owner'] = owner
                else:
                    row_errors.append(ImportError(
                        row=row_num, field='owner_username_or_phone',
                        message=f'User "{owner_ref}" not found.'
                    ))
            
            if row_errors:
                errors.extend(row_errors)
            else:
                validated_data['_row_num'] = row_num
                validated_rows.append(validated_data)
        
        return validated_rows, errors

    @classmethod
    def _build_category_lookup(cls) -> Dict[str, Category]:
        """
        Build lookup dict for categories by full path.
        
        Returns:
            Dict mapping lowercase full_path to Category
        """
        categories = {}
        for cat in Category.objects.filter(is_active=True):
            categories[cat.full_path.lower()] = cat
            # Also allow lookup by name only for root categories
            if cat.parent_id is None:
                categories[cat.name.lower()] = cat
        return categories
    
    @classmethod
    def _build_user_lookup(cls) -> Dict[str, User]:
        """
        Build lookup dict for users by username or phone.
        
        Returns:
            Dict mapping lowercase username/phone to User
        """
        users = {}
        for user in User.objects.filter(is_active=True):
            users[user.username.lower()] = user
            if user.phone:
                users[user.phone.lower()] = user
        return users
    
    @classmethod
    def _get_audit_data(cls, item: Item) -> dict:
        """Extract audit-safe data from item."""
        return {
            'name': item.name,
            'sku': item.sku,
            'barcode': item.barcode,
            'category_id': item.category_id,
            'unit': item.unit,
            'description': item.description,
            'location_id': item.location_id,
            'owner_id': item.owner_id,
            'type': item.type,
            'status': item.status,
            'quantity': item.quantity,
            'is_active': item.is_active,
        }
    
    @classmethod
    @transaction.atomic
    def import_items(
        cls,
        file_content: bytes,
        file_name: str,
        operator,
        ip: Optional[str] = None,
        user_agent: str = '',
        allow_partial: bool = True
    ) -> ImportResult:
        """
        Import items from CSV/XLSX file.
        
        Uses SKU as natural key for upsert:
        - If item with SKU exists: update it
        - If item with SKU doesn't exist: create it
        
        Args:
            file_content: Raw bytes of file
            file_name: Original filename (used to determine format)
            operator: User performing the import
            ip: Client IP address
            user_agent: Client user agent
            allow_partial: If True, import valid rows even if some have errors
            
        Returns:
            ImportResult with counts and errors
            
        Raises:
            BusinessConflictError: If file is invalid or too large
        """
        result = ImportResult()
        
        # Validate file size
        cls.validate_file_size(len(file_content))
        
        # Parse file based on extension
        file_ext = file_name.lower().split('.')[-1] if '.' in file_name else ''
        if file_ext == 'csv':
            headers, rows = cls.parse_csv(file_content)
        elif file_ext in ['xlsx', 'xls']:
            headers, rows = cls.parse_xlsx(file_content)
        else:
            raise BusinessConflictError(
                code='INVALID_FILE_TYPE',
                message='File must be CSV or XLSX format.',
                details={'extension': file_ext}
            )
        
        # Validate row count
        cls.validate_row_count(len(rows))
        
        # Validate headers
        header_errors = cls.validate_headers(headers)
        if header_errors:
            result.errors.extend(header_errors)
            result.error_count = len(header_errors)
            return result
        
        # Validate all rows
        validated_rows, validation_errors = cls.validate_rows(rows, headers)
        result.errors.extend(validation_errors)
        result.error_count = len(validation_errors)
        
        # If not allowing partial and there are errors, return early
        if not allow_partial and validation_errors:
            return result
        
        # Process validated rows
        existing_skus = set(
            Item.objects.filter(
                sku__in=[r['sku'] for r in validated_rows]
            ).values_list('sku', flat=True)
        )
        
        for row_data in validated_rows:
            try:
                sku = row_data['sku']
                is_update = sku in existing_skus
                
                if is_update:
                    # Update existing item
                    item = Item.objects.select_for_update().get(sku=sku)
                    before = cls._get_audit_data(item)
                    
                    item.name = row_data['name']
                    item.location = row_data['location']
                    
                    if 'barcode' in row_data:
                        item.barcode = row_data['barcode']
                    if 'unit' in row_data:
                        item.unit = row_data['unit']
                    if 'description' in row_data:
                        item.description = row_data['description']
                    if 'category' in row_data:
                        item.category = row_data['category']
                    if 'owner' in row_data:
                        item.owner = row_data['owner']
                    
                    # Update quantity for consumable items
                    if item.type == ItemType.CONSUMABLE and 'quantity' in row_data:
                        item.quantity = row_data['quantity']
                        item.status = ItemStatus.IN_STOCK if item.quantity > 0 else ItemStatus.OUT_OF_STOCK
                    
                    item.save()
                    
                    after = cls._get_audit_data(item)
                    if before != after:
                        AuditLog.objects.create(
                            action=AuditAction.UPDATE,
                            target_type=TargetType.ITEM,
                            target_id=item.id,
                            diff={'before': before, 'after': after},
                            operator=operator,
                            operator_role=getattr(operator, 'role', ''),
                            ip=ip,
                            user_agent=user_agent,
                        )
                    
                    result.update_count += 1
                else:
                    # Create new item
                    item_type = row_data['type']
                    quantity = row_data.get('quantity')
                    
                    # Determine initial status
                    if item_type == ItemType.REUSABLE:
                        status = ItemStatus.AVAILABLE
                        quantity = None
                    else:
                        quantity = quantity if quantity is not None else 0
                        status = ItemStatus.IN_STOCK if quantity > 0 else ItemStatus.OUT_OF_STOCK
                    
                    item = Item.objects.create(
                        name=row_data['name'],
                        sku=sku,
                        barcode=row_data.get('barcode'),
                        category=row_data.get('category'),
                        unit=row_data.get('unit', ''),
                        description=row_data.get('description', ''),
                        location=row_data['location'],
                        owner=row_data.get('owner', operator),
                        type=item_type,
                        status=status,
                        quantity=quantity,
                        is_active=True,
                    )
                    
                    AuditLog.objects.create(
                        action=AuditAction.CREATE,
                        target_type=TargetType.ITEM,
                        target_id=item.id,
                        diff={'before': None, 'after': cls._get_audit_data(item)},
                        operator=operator,
                        operator_role=getattr(operator, 'role', ''),
                        ip=ip,
                        user_agent=user_agent,
                    )
                    
                    result.create_count += 1
                
                result.success_count += 1
                
            except Exception as e:
                row_num = row_data.get('_row_num', 0)
                result.errors.append(ImportError(
                    row=row_num,
                    field='',
                    message=f'Failed to import row: {str(e)}'
                ))
                result.error_count += 1
        
        # Create audit log for the import operation itself
        AuditLog.objects.create(
            action=AuditAction.IMPORT,
            target_type=TargetType.ITEM,
            target_id=0,  # No specific target
            diff={
                'file_name': file_name,
                'total_rows': len(rows),
                'success_count': result.success_count,
                'create_count': result.create_count,
                'update_count': result.update_count,
                'error_count': result.error_count,
            },
            operator=operator,
            operator_role=getattr(operator, 'role', ''),
            ip=ip,
            user_agent=user_agent,
        )
        
        return result
    
    @classmethod
    def get_import_template_headers(cls) -> List[str]:
        """
        Get headers for import template.
        
        Returns:
            List of column headers
        """
        return [
            'name',
            'sku',
            'type',
            'warehouse_code',
            'location_code',
            'barcode',
            'quantity',
            'unit',
            'description',
            'category_path',
            'owner_username_or_phone',
        ]
    
    @classmethod
    def generate_import_template_csv(cls) -> bytes:
        """
        Generate a CSV template for import.
        
        Returns:
            CSV file content as bytes
        """
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Write headers
        writer.writerow(cls.get_import_template_headers())
        
        # Write example row
        writer.writerow([
            'Example Item',
            'SKU-001',
            'reusable',
            'WH-001',
            'A-01',
            '',
            '',
            'pcs',
            'Example description',
            'Electronics/Computers',
            '',
        ])
        
        return output.getvalue().encode('utf-8-sig')
    
    @classmethod
    def generate_import_template_xlsx(cls) -> bytes:
        """
        Generate an XLSX template for import.
        
        Returns:
            XLSX file content as bytes
        """
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        
        wb = Workbook()
        ws = wb.active
        ws.title = 'Items'
        
        headers = cls.get_import_template_headers()
        
        # Write headers with styling
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
        
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        
        # Write example row
        example_row = [
            'Example Item',
            'SKU-001',
            'reusable',
            'WH-001',
            'A-01',
            '',
            '',
            'pcs',
            'Example description',
            'Electronics/Computers',
            '',
        ]
        for col, value in enumerate(example_row, start=1):
            ws.cell(row=2, column=col, value=value)
        
        # Adjust column widths
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 20
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()
