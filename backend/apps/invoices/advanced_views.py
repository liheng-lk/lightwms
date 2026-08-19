"""Advanced invoice filtering and export endpoints."""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.response import Response
from rest_framework.views import APIView

from core.pagination import StandardPagination
from core.permissions import IsSuperAdmin
from .global_scope_views import InvoiceListCreateView as BaseInvoiceListCreateView
from .hardened_views import (
    InvoiceExportView as BaseInvoiceExportView,
    _export_files_to_safe_zip,
)
from .serializers import InvoiceSerializer
from .services import InvoiceService
from core.exceptions import BusinessConflictError


def _parse_date(value: str | None, field: str) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f'{field} 必须是 YYYY-MM-DD 格式') from exc


def _filtered_invoices(params):
    uploader = params.get('uploader') or params.get('created_by')
    try:
        uploader_id = int(uploader) if uploader else None
    except (TypeError, ValueError) as exc:
        raise ValueError('uploader 必须是有效的用户 ID') from exc

    date_from = _parse_date(params.get('date_from'), 'date_from')
    date_to = _parse_date(params.get('date_to'), 'date_to')
    if date_from and date_to and date_from > date_to:
        raise ValueError('开始日期不能晚于结束日期')

    item = params.get('item')
    try:
        item_id = int(item) if item else None
    except (TypeError, ValueError) as exc:
        raise ValueError('item 必须是有效的物品 ID') from exc

    invoices = InvoiceService.list(
        category=params.get('category'),
        status=params.get('status'),
        registered_month=params.get('month'),
        item_id=item_id,
        created_by_id=uploader_id,
        search=params.get('search'),
    )
    if date_from:
        invoices = invoices.filter(created_at__date__gte=date_from)
    if date_to:
        invoices = invoices.filter(created_at__date__lte=date_to)
    return invoices


def _validation_response(exc: ValueError) -> Response:
    return Response(
        {'code': 'INVALID_FILTER', 'message': str(exc), 'details': None},
        status=status.HTTP_400_BAD_REQUEST,
    )


class InvoiceListCreateView(BaseInvoiceListCreateView):
    """Invoice list with uploader and custom date-range filtering."""

    def get(self, request):
        try:
            invoices = _filtered_invoices(request.query_params)
        except ValueError as exc:
            return _validation_response(exc)

        paginator = StandardPagination()
        page = paginator.paginate_queryset(invoices, request)
        if page is not None:
            return paginator.get_paginated_response(InvoiceSerializer(page, many=True).data)
        return Response(InvoiceSerializer(invoices, many=True).data, status=status.HTTP_200_OK)


class InvoiceExportView(BaseInvoiceExportView):
    """ZIP invoice-file export with uploader and custom date range support."""

    def get(self, request):
        try:
            invoices = _filtered_invoices(request.query_params)
        except ValueError as exc:
            return _validation_response(exc)

        try:
            content = _export_files_to_safe_zip(invoices)
        except BusinessConflictError as exc:
            return Response(
                {'code': exc.code, 'message': exc.message, 'details': exc.details},
                status=status.HTTP_409_CONFLICT,
            )

        start = request.query_params.get('date_from') or 'all'
        end = request.query_params.get('date_to') or 'all'
        response = HttpResponse(content, content_type='application/zip')
        response['Content-Disposition'] = (
            f'attachment; filename="invoices_{start}_{end}.zip"'
        )
        return response


def _reimbursement_excel(invoices) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '报销清单'
    sheet.freeze_panes = 'A2'

    headers = [
        '序号', '发票名称', '金额', '类别', '类型', '项目', '报销状态',
        '关联物品', '登记月份', '上传者', '上传日期', '原文件名',
    ]
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)

    total = Decimal('0')
    row = 2
    for index, invoice in enumerate(invoices, 1):
        amount = Decimal(invoice.amount)
        total += amount
        values = [
            index,
            invoice.name,
            float(amount),
            invoice.category,
            invoice.custom_type,
            invoice.project,
            invoice.get_status_display(),
            invoice.item.name if invoice.item else '',
            invoice.registered_month,
            invoice.created_by.username if invoice.created_by else '',
            invoice.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            invoice.original_filename,
        ]
        for column, value in enumerate(values, 1):
            sheet.cell(row=row, column=column, value=value)
        row += 1

    sheet.cell(row=row, column=1, value='合计').font = Font(bold=True)
    sheet.cell(row=row, column=3, value=float(total)).font = Font(bold=True)
    sheet.auto_filter.ref = f'A1:L{max(1, row - 1)}'

    widths = [8, 24, 14, 16, 16, 20, 14, 22, 14, 18, 20, 28]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class ReimbursementExcelExportView(APIView):
    """Export the current invoice filter as an Excel reimbursement list."""

    permission_classes = [IsSuperAdmin]

    def get(self, request):
        try:
            invoices = _filtered_invoices(request.query_params)
        except ValueError as exc:
            return _validation_response(exc)

        content = _reimbursement_excel(invoices)
        start = request.query_params.get('date_from') or 'all'
        end = request.query_params.get('date_to') or 'all'
        response = HttpResponse(
            content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = (
            f'attachment; filename="reimbursement_{start}_{end}.xlsx"'
        )
        return response
