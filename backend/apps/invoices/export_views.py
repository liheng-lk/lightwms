"""Invoice export endpoints with shared filters and explicit file formats."""
from __future__ import annotations

import io
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.response import Response
from rest_framework.views import APIView

from core.exceptions import BusinessConflictError
from core.permissions import IsItemManagerOrSuperAdmin, IsSuperAdmin

from .filters import InvoiceFilterError, get_filtered_invoices
from .hardened_views import (
    InvoiceExportView as HardenedInvoiceExportView,
    _export_files_to_safe_zip,
    _normalise_ids,
)
from .models import Invoice


def _filter_error_response(exc: InvoiceFilterError):
    return Response(
        {'code': 'INVALID_FILTER', 'message': str(exc), 'details': None},
        status=status.HTTP_400_BAD_REQUEST,
    )


def _selected_invoices(raw_ids):
    ids = _normalise_ids(raw_ids)
    if ids is None:
        return None, Response(
            {'code': 'INVALID_REQUEST', 'message': 'ids must be a non-empty list', 'details': None},
            status=status.HTTP_400_BAD_REQUEST,
        )

    invoices = list(
        Invoice.objects.select_related('item', 'created_by')
        .filter(id__in=ids)
        .order_by('id')
    )
    if len(invoices) != len(ids):
        found = {invoice.id for invoice in invoices}
        missing = [invoice_id for invoice_id in ids if invoice_id not in found]
        return None, Response(
            {
                'code': 'NOT_FOUND',
                'message': 'One or more invoices were not found.',
                'details': {'ids': missing},
            },
            status=status.HTTP_404_NOT_FOUND,
        )
    return invoices, None


def _export_reimbursement_workbook(invoices) -> bytes:
    """Build a finance-friendly reimbursement workbook from invoice rows."""
    rows = list(invoices)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '报销清单'
    sheet.freeze_panes = 'A2'

    headers = [
        '序号',
        '发票名称',
        '金额（元）',
        '类别',
        '类型',
        '项目',
        '报销状态',
        '关联物品',
        '登记月份',
        '上传者',
        '上传时间',
        '原始文件名',
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    total = Decimal('0.00')
    for index, invoice in enumerate(rows, start=1):
        total += invoice.amount
        sheet.append([
            index,
            invoice.name,
            float(invoice.amount),
            invoice.category,
            invoice.custom_type,
            invoice.project,
            invoice.get_status_display(),
            invoice.item.name if invoice.item else '',
            invoice.registered_month,
            invoice.created_by.username if invoice.created_by else '',
            invoice.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            invoice.original_filename,
        ])
        sheet.cell(row=index + 1, column=3).number_format = '#,##0.00'

    total_row = len(rows) + 2
    sheet.cell(row=total_row, column=2, value='合计')
    sheet.cell(row=total_row, column=2).font = Font(bold=True)
    sheet.cell(row=total_row, column=3, value=float(total))
    sheet.cell(row=total_row, column=3).font = Font(bold=True)
    sheet.cell(row=total_row, column=3).number_format = '#,##0.00'

    if rows:
        sheet.auto_filter.ref = f'A1:L{len(rows) + 1}'

    widths = [8, 28, 14, 16, 16, 22, 14, 24, 12, 18, 20, 30]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.read()


class InvoiceExportView(HardenedInvoiceExportView):
    """Export original invoice files as ZIP using the same filters as the list."""

    permission_classes = [IsSuperAdmin]

    def get(self, request):
        try:
            invoices = get_filtered_invoices(request.query_params)
            content = _export_files_to_safe_zip(invoices)
        except InvoiceFilterError as exc:
            return _filter_error_response(exc)
        except BusinessConflictError as exc:
            return Response(
                {'code': exc.code, 'message': exc.message, 'details': exc.details},
                status=status.HTTP_409_CONFLICT,
            )

        date_from = request.query_params.get('date_from') or 'all'
        date_to = request.query_params.get('date_to') or 'all'
        response = HttpResponse(content, content_type='application/zip')
        response['Content-Disposition'] = (
            f'attachment; filename="invoices_{date_from}_{date_to}.zip"'
        )
        return response


class InvoiceReimbursementExportView(APIView):
    """Export a reimbursement list as XLSX, filtered or selected by invoice IDs."""

    permission_classes = [IsSuperAdmin]

    def _response(self, invoices, filename: str):
        content = _export_reimbursement_workbook(invoices)
        response = HttpResponse(
            content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    def get(self, request):
        try:
            invoices = get_filtered_invoices(request.query_params)
        except InvoiceFilterError as exc:
            return _filter_error_response(exc)

        date_from = request.query_params.get('date_from') or 'all'
        date_to = request.query_params.get('date_to') or 'all'
        return self._response(
            invoices,
            f'reimbursement_list_{date_from}_{date_to}.xlsx',
        )

    def post(self, request):
        invoices, error_response = _selected_invoices(request.data.get('ids'))
        if error_response is not None:
            return error_response
        return self._response(
            invoices,
            f'reimbursement_selected_{len(invoices)}.xlsx',
        )


class InvoiceUploadersView(APIView):
    """Return uploaders that actually have invoice records."""

    permission_classes = [IsItemManagerOrSuperAdmin]

    def get(self, request):
        rows = (
            Invoice.objects.values('created_by_id', 'created_by__username')
            .distinct()
            .order_by('created_by__username', 'created_by_id')
        )
        return Response(
            [
                {'id': row['created_by_id'], 'username': row['created_by__username']}
                for row in rows
            ],
            status=status.HTTP_200_OK,
        )
